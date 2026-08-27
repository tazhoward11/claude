"""Merge the two Apollo runs for the ethnic-owned business list into one clean sheet.

Selection follows the standing rules: max 4 contacts per company, decision makers
(owner/founder/C-level/president/VP/director) and the highest marketing person
first, junior marketing only as an execution contact, Austin corridor only.
"""
import csv
from collections import defaultdict, OrderedDict

import openpyxl
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

# Later files win, so the first-name-format re-run overwrites the blank rows.
SOURCES = ["ethnic_leads.csv", "ethnic_first_fmt.csv"]

# Curated picks, in outreach priority order. Anyone not listed here was dropped as
# not a buying decision maker (HR managers, housekeeping, data entry, associate
# media directors, service advisors).
PICKS = OrderedDict([
    ("Third Ear", ["Alejandro Ru***s", "Cailin Bu***a", "Serge Fl***s", "Charles Ne***r"]),
    ("CarbonBetter", ["Tri Vo", "Natalie Butler", "Matt Hendren", "Laura Al***r"]),
    ("SmartTouch Interactive", ["Robert Cowes", "Aaron Fichera", "Tanner Ro***s", "Henry Me***o"]),
    ("Journeyman Group", ["Kaleigh Laduca", "Juanita Kniaz", "David Gr***k", "Cannon Ki***e"]),
    ("The Personnel Store", ["Scott Cunningham", "Travis Cu***m", "Mac Cunningham"]),
    ("KLP Construction Supply", ["Nathali Parker", "Karen Rogers", "Brandon Sc***h"]),
    ("Express Commercial Cleaning", ["Evelyn Tavernier", "Dan Tavernier"]),
    ("Moore Clean", ["Jose Moore", "Scott Oliver"]),
    ("CyberTex", ["Monique Johnson", "Monika Whitaker"]),
    ("Austin Underground", ["Marena Gallo", "Kenna Tolman"]),
    ("Building Team Solutions", ["Britanie Olvera"]),
    ("Austin Staffing", ["Chris Guaydacan"]),
    ("Modular Installation Services", ["Monica Gould"]),
    ("Neighborhood Plumbing and Drain", ["Heberto Alanis"]),
])

NO_APOLLO = [
    ("AvantGarde", "No Austin org in Apollo (only a Munich, Germany company by that name)"),
    ("Environments Plus", "Only the Los Angeles firm is in Apollo, not an Austin one"),
    ("Michael's Insurance Group", "No Apollo record at all"),
    ("LOC Consultants", "Org exists (9 employees, Austin) but Apollo has no domain, so no email is inferable"),
    ("Pro Tech Construction", "4 employees in Apollo, none in the Austin corridor"),
    ("Alpha One Motors", "Only a Client Advisor and a Service Manager - no decision maker"),
    ("TRI Recycling", "Only a Data Entry Specialist"),
    ("Best Choice Mobile Notary", "Only a Notary/Trainer"),
    ("This Lil' Dog of Mine", "Zero people in Apollo"),
    ("Total Leadership Ventures", "Only a Career Coach"),
    ("Azarmehr Law Group", "Only attorneys and paralegals - no exec or marketing contact"),
]

CONFIDENCE = {
    "verified": "Verified by Apollo",
    "predicted (not verified)": "Pattern-based - verify before sending",
    "unavailable": "",
}


def load():
    by_key = {}
    for path in SOURCES:
        try:
            fh = open(path, newline="")
        except FileNotFoundError:
            continue
        with fh:
            for r in csv.DictReader(fh):
                key = (r["company"], r["name"])
                # Keep the richer row: never let a later blank email clobber a real one.
                prev = by_key.get(key)
                if prev and prev.get("email") and not r.get("email"):
                    continue
                by_key[key] = r
    return by_key


def main():
    by_key = load()
    rows, manual = [], []

    for company, names in PICKS.items():
        for name in names:
            r = by_key.get((company, name))
            if r is None:
                manual.append([company, name, "", "Not found in the merged Apollo output"])
                continue
            email, status = r.get("email", ""), r.get("email_status", "")
            if not email:
                reason = ("Apollo has no email on file for this person"
                          if status == "unavailable"
                          else "Surname is masked by Apollo and this domain does not use a "
                               "first-name-only format, so the address cannot be inferred")
                manual.append([company, name, r.get("title", ""), reason])
                continue
            rows.append([company, name, r.get("title", ""), email, CONFIDENCE.get(status, status)])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Contacts"

    head = Font(bold=True, color="FFFFFF")
    fill = openpyxl.styles.PatternFill("solid", fgColor="1F3864")
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def sheet(ws, headers, data, widths):
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill = head, fill
            c.alignment = Alignment(vertical="center")
        for row in data:
            ws.append(row)
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=len(headers)):
            for c in row:
                c.border = border
                c.alignment = Alignment(vertical="top", wrap_text=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    sheet(ws, ["Company", "Name", "Title", "Email", "Confidence"], rows, [30, 24, 42, 40, 30])

    ws2 = wb.create_sheet("Needs Manual Research")
    sheet(ws2, ["Company", "Name", "Title", "Why"],
          manual + [[c, "", "", why] for c, why in NO_APOLLO], [30, 24, 40, 62])

    wb.save("ethnic_owned_contacts.xlsx")
    print(f"Contacts: {len(rows)} across {len({r[0] for r in rows})} companies")
    print(f"Manual:   {len(manual)} people + {len(NO_APOLLO)} companies with no usable Apollo record")
    ver = sum(1 for r in rows if r[4].startswith("Verified"))
    print(f"          {ver} verified, {len(rows) - ver} pattern-based")


if __name__ == "__main__":
    main()
